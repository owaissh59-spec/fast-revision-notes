#!/usr/bin/env python3
"""
docxkit.py -- Word-document building blocks for the surgical-tray notes.

Design goals
------------
* Legal page (8.5 x 14 in) with 1.27 cm margins on all four sides.
* Pure monochrome: distinction comes from weight, rules, borders and grey
  shading only -- nothing depends on colour.
* A real, updatable Table of Contents (TOC field) plus explicit bookmarks
  on every heading so the TOC and any cross-references can be refreshed
  in Word with Ctrl+A -> F9.
* Images sit in the RIGHT column of a two-column layout table so each
  figure sits beside the text it belongs to.
"""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

# ----------------------------------------------------------------- constants
PAGE_W = Cm(21.59)          # 8.5 in
PAGE_H = Cm(35.56)          # 14 in  (US Legal)
MARGIN = Cm(1.27)           # 0.5 in "narrow"
BODY_W = Cm(21.59 - 2 * 1.27)   # 19.05 cm usable

# grey levels used for shading (hex, monochrome-safe)
G_HEAD = "1A1A1A"           # near-black banner
G_SUB = "555555"
G_BOX = "F0F0F0"            # light box fill
G_BOX2 = "E4E4E4"
G_BAND = "F6F6F6"           # table row banding
G_TABHEAD = "D9D9D9"
G_RULE = "808080"

FONT_BODY = "Georgia"
FONT_HEAD = "Calibri"
FONT_MONO = "Consolas"

_bookmark_id = [1000]


# ------------------------------------------------------------------ xml utils
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def _sub(parent, tag, **attrs):
    e = _el(tag, **attrs)
    parent.append(e)
    return e


def shade(obj, fill):
    """Apply cell / paragraph shading."""
    pr = obj._tc.get_or_add_tcPr() if hasattr(obj, "_tc") else \
        obj._p.get_or_add_pPr()
    for old in pr.findall(qn("w:shd")):
        pr.remove(old)
    pr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto",
                              "w:fill": fill}))


def cell_margins(cell, top=60, start=90, bottom=60, end=90):
    """Set cell padding in twentieths of a point."""
    tcPr = cell._tc.get_or_add_tcPr()
    m = _sub(tcPr, "w:tcMar")
    for tag, val in (("w:top", top), ("w:start", start),
                     ("w:bottom", bottom), ("w:end", end)):
        _sub(m, tag, **{"w:w": val, "w:type": "dxa"})


def set_borders(obj, edges, sz=6, val="single", color=G_RULE, space=0):
    """
    Borders for a table, cell or paragraph.
    edges: dict like {"top": {...}} or a list of edge names using defaults.
    """
    if isinstance(edges, (list, tuple, set)):
        edges = {e: {} for e in edges}
    if hasattr(obj, "_tbl"):
        pr = obj._tbl.tblPr
        tag = "w:tblBorders"
    elif hasattr(obj, "_tc"):
        pr = obj._tc.get_or_add_tcPr()
        tag = "w:tcBorders"
    else:
        pr = obj._p.get_or_add_pPr()
        tag = "w:pBdr"
    for old in pr.findall(qn(tag)):
        pr.remove(old)
    b = _el(tag)
    order = ["top", "left", "start", "bottom", "right", "end",
             "insideH", "insideV"]
    for name in order:
        if name not in edges:
            continue
        spec = edges[name]
        b.append(_el(f"w:{name}", **{
            "w:val": spec.get("val", val),
            "w:sz": spec.get("sz", sz),
            "w:space": spec.get("space", space),
            "w:color": spec.get("color", color),
        }))
    pr.append(b)


def no_borders(table):
    set_borders(table, {e: {"val": "nil"} for e in
                        ("top", "left", "bottom", "right", "insideH",
                         "insideV")})


def keep_together(p, with_next=True):
    pPr = p._p.get_or_add_pPr()
    _sub(pPr, "w:keepLines")
    if with_next:
        _sub(pPr, "w:keepNext")


def widow_control(p, on=True):
    pPr = p._p.get_or_add_pPr()
    _sub(pPr, "w:widowControl", **{"w:val": "1" if on else "0"})


def spacing(p, before=None, after=None, line=None, rule="auto"):
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:spacing")):
        pPr.remove(old)
    attrs = {}
    if before is not None:
        attrs["w:before"] = int(before * 20)
    if after is not None:
        attrs["w:after"] = int(after * 20)
    if line is not None:
        attrs["w:line"] = int(line * 240)
        attrs["w:lineRule"] = rule
    pPr.append(_el("w:spacing", **attrs))


def indent(p, left=None, right=None, first=None, hanging=None):
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:ind")):
        pPr.remove(old)
    attrs = {}
    if left is not None:
        attrs["w:start"] = int(left * 567)      # cm -> twips
    if right is not None:
        attrs["w:end"] = int(right * 567)
    if first is not None:
        attrs["w:firstLine"] = int(first * 567)
    if hanging is not None:
        attrs["w:hanging"] = int(hanging * 567)
    pPr.append(_el("w:ind", **attrs))


# ---------------------------------------------------------------- field codes
def add_field(p, instr, placeholder="", bold=False, size=None, font=None):
    """Insert a complex Word field (TOC, PAGE, NUMPAGES, REF ...)."""
    r1 = p.add_run()
    r1._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    r2 = p.add_run()
    it = _el("w:instrText", **{"xml:space": "preserve"})
    it.text = instr
    r2._r.append(it)
    r3 = p.add_run()
    r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    r4 = p.add_run(placeholder)
    if bold:
        r4.bold = True
    if size:
        r4.font.size = Pt(size)
    if font:
        r4.font.name = font
    r5 = p.add_run()
    r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    return r4


def bookmark(p, name):
    """Wrap a paragraph in a bookmark so REF / TOC can target it."""
    _bookmark_id[0] += 1
    bid = _bookmark_id[0]
    safe = "".join(ch if (ch.isalnum() or ch == "_") else "_" for ch in name)
    start = _el("w:bookmarkStart", **{"w:id": bid, "w:name": safe})
    end = _el("w:bookmarkEnd", **{"w:id": bid})
    p._p.insert(0, start)
    p._p.append(end)
    return safe


def update_fields_on_open(doc):
    """Tell Word to refresh all fields (so the TOC fills in) when opened."""
    settings = doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    settings.append(_el("w:updateFields", **{"w:val": "true"}))


# ------------------------------------------------------------------- document
def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = PAGE_W
    sec.page_height = PAGE_H
    sec.left_margin = sec.right_margin = MARGIN
    sec.top_margin = sec.bottom_margin = MARGIN
    sec.header_distance = Cm(0.6)
    sec.footer_distance = Cm(0.6)
    _build_styles(doc)
    return doc


def _style(doc, name, base=None, font=FONT_BODY, size=10.5, bold=False,
           italic=False, color="000000", space_before=0, space_after=0,
           line=1.12, align=None, caps=False, spacing_pts=None,
           keep_next=False, outline=None):
    styles = doc.styles
    try:
        st = styles[name]
    except KeyError:
        from docx.enum.style import WD_STYLE_TYPE
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if base:
            st.base_style = styles[base]
    f = st.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = RGBColor.from_string(color)
    if caps:
        f.all_caps = True
    if spacing_pts is not None:
        rPr = st.element.get_or_add_rPr()
        rPr.append(_el("w:spacing", **{"w:val": int(spacing_pts * 20)}))
    # east-asian + complex-script font names
    rPr = st.element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = _sub(rPr, "w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), font)
    pf = st.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    pPr = st.element.get_or_add_pPr()
    if keep_next:
        _sub(pPr, "w:keepNext")
        _sub(pPr, "w:keepLines")
    if outline is not None:
        _sub(pPr, "w:outlineLvl", **{"w:val": outline})
    return st


def _build_styles(doc):
    # ---- base body text
    base = doc.styles["Normal"]
    base.font.name = FONT_BODY
    base.font.size = Pt(10.5)
    rPr = base.element.get_or_add_rPr()
    rf = _sub(rPr, "w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), FONT_BODY)
    base.paragraph_format.space_after = Pt(3)
    base.paragraph_format.line_spacing = 1.12

    # ---- headings (map to Heading 1..4 so the TOC picks them up)
    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT_HEAD
    h1.font.size = Pt(19)
    h1.font.bold = True
    h1.font.all_caps = True
    h1.font.color.rgb = RGBColor.from_string("FFFFFF")
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT_HEAD
    h2.font.size = Pt(14.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string("000000")
    h2.paragraph_format.space_before = Pt(15)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    h3.font.name = FONT_HEAD
    h3.font.size = Pt(11.8)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string("1A1A1A")
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True

    h4 = doc.styles["Heading 4"]
    h4.font.name = FONT_HEAD
    h4.font.size = Pt(10.5)
    h4.font.bold = True
    h4.font.italic = True
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(2)
    h4.paragraph_format.keep_with_next = True

    # ---- custom styles
    _style(doc, "TrayLead", font=FONT_BODY, size=10.5, line=1.14,
           space_after=4)
    _style(doc, "TinyLabel", font=FONT_HEAD, size=8, bold=True, caps=True,
           color="333333", space_after=1, spacing_pts=0.6, keep_next=True)
    _style(doc, "Bullet", font=FONT_BODY, size=10.2, line=1.10, space_after=2)
    _style(doc, "SubBullet", font=FONT_BODY, size=9.8, line=1.08, space_after=1)
    _style(doc, "BoxHead", font=FONT_HEAD, size=9.5, bold=True, caps=True,
           color="000000", space_after=2, spacing_pts=0.5, keep_next=True)
    _style(doc, "BoxBody", font=FONT_BODY, size=9.6, line=1.10, space_after=2)
    _style(doc, "TabHead", font=FONT_HEAD, size=9.2, bold=True, line=1.05,
           space_before=1, space_after=1, keep_next=True)
    _style(doc, "TabCell", font=FONT_BODY, size=9.4, line=1.06,
           space_before=1, space_after=1)
    _style(doc, "TabCellSm", font=FONT_BODY, size=8.9, line=1.05,
           space_before=1, space_after=1)
    _style(doc, "FigCaption", font=FONT_HEAD, size=8.4, italic=True,
           color="333333", line=1.05, space_before=2, space_after=6,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _style(doc, "TOCHead", font=FONT_HEAD, size=20, bold=True, caps=True,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4, spacing_pts=1.2)
    _style(doc, "CoverTitle", font=FONT_HEAD, size=34, bold=True,
           align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, space_after=0)
    _style(doc, "CoverSub", font=FONT_HEAD, size=13, bold=False,
           align=WD_ALIGN_PARAGRAPH.CENTER, color="333333", space_after=0)
    _style(doc, "Mnemonic", font=FONT_MONO, size=9.8, bold=True, line=1.1,
           space_after=2)
    _style(doc, "QBank", font=FONT_BODY, size=9.8, line=1.10, space_after=3)


# --------------------------------------------------------------- header/footer
def set_header_footer(doc, left_text, right_text="Instrument Trays"):
    sec = doc.sections[0]

    hp = sec.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.add_tab_stop(BODY_W, WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(left_text)
    r.font.name = FONT_HEAD
    r.font.size = Pt(8.2)
    r.font.bold = True
    r.font.all_caps = True
    r.font.color.rgb = RGBColor.from_string("444444")
    r2 = hp.add_run("\t" + right_text)
    r2.font.name = FONT_HEAD
    r2.font.size = Pt(8.2)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor.from_string("666666")
    set_borders(hp, {"bottom": {"sz": 6, "color": "AAAAAA", "space": 2}})
    spacing(hp, after=4)

    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = fp.add_run("— ")
    add_field(fp, "PAGE", "1", bold=True, size=9.5, font=FONT_HEAD)
    r4 = fp.add_run(" —")
    for rr in (r3, r4):
        rr.font.name = FONT_HEAD
        rr.font.size = Pt(9.5)
        rr.font.color.rgb = RGBColor.from_string("555555")
    set_borders(fp, {"top": {"sz": 6, "color": "AAAAAA", "space": 2}})
    spacing(fp, before=3)


# ----------------------------------------------------------------- primitives
def para(doc, text="", style="Normal", **kw):
    p = doc.add_paragraph(style=style)
    if text:
        _rich(p, text)
    if kw:
        spacing(p, **{k: v for k, v in kw.items()
                      if k in ("before", "after", "line")})
    return p


def _rich(p, text, base_bold=False, base_italic=False):
    """
    Mini-markup:  **bold**   //italic//   __underline__   `mono`
    Segments may not nest.
    """
    import re
    tokens = re.split(r"(\*\*.+?\*\*|//.+?//|__.+?__|`.+?`)", text)
    for t in tokens:
        if not t:
            continue
        if t.startswith("**") and t.endswith("**") and len(t) > 4:
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("//") and t.endswith("//") and len(t) > 4:
            r = p.add_run(t[2:-2]); r.italic = True
        elif t.startswith("__") and t.endswith("__") and len(t) > 4:
            r = p.add_run(t[2:-2]); r.underline = True
        elif t.startswith("`") and t.endswith("`") and len(t) > 2:
            r = p.add_run(t[1:-1]); r.font.name = FONT_MONO
            r.font.size = Pt(9.4)
        else:
            r = p.add_run(t)
        if base_bold:
            r.bold = True
        if base_italic:
            r.italic = True
    return p


def bullet(doc, text, level=0, marker=None, style=None):
    style = style or ("Bullet" if level == 0 else "SubBullet")
    mk = marker if marker is not None else ("▪" if level == 0 else "–")
    p = doc.add_paragraph(style=style)
    indent(p, left=0.42 + 0.42 * level, hanging=0.42)
    r = p.add_run(mk + "  ")
    r.font.name = FONT_HEAD
    r.bold = level == 0
    _rich(p, text)
    return p


def numbered(doc, n, text, level=0):
    p = doc.add_paragraph(style="Bullet" if level == 0 else "SubBullet")
    indent(p, left=0.55 + 0.42 * level, hanging=0.55)
    r = p.add_run(f"{n}.  ")
    r.font.name = FONT_HEAD
    r.bold = True
    _rich(p, text)
    return p


def rule(doc, sz=8, color="000000", before=2, after=6):
    p = doc.add_paragraph()
    spacing(p, before=before, after=after)
    p.paragraph_format.line_spacing = 1
    for r in p.runs:
        r.text = ""
    set_borders(p, {"bottom": {"sz": sz, "color": color, "space": 1}})
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p



# ==========================================================================
#  High-level components
# ==========================================================================
def part_banner(doc, number, title, subtitle=None):
    """Full-width reversed-out (black) part banner. Bookmarked for the TOC."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c = t.rows[0].cells[0]
    c.width = BODY_W
    shade(c, G_HEAD)
    cell_margins(c, top=150, bottom=150, start=220, end=220)
    no_borders(t)

    # Kicker line ("PART I") is a plain paragraph, NOT the heading, so the
    # TOC entry contains only the part title.
    p0 = c.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    spacing(p0, before=0, after=2)
    p0.paragraph_format.line_spacing = 1
    keep_together(p0)
    r = p0.add_run(f"PART {number}")
    r.font.name = FONT_HEAD
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("CCCCCC")
    r.font.all_caps = True
    rPr = r._r.get_or_add_rPr()
    rPr.append(_el("w:spacing", **{"w:val": 50}))

    p = c.add_paragraph(style=doc.styles["Heading 1"])
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    spacing(p, before=0, after=0)
    r2 = p.add_run(title)
    r2.font.name = FONT_HEAD
    r2.font.size = Pt(19)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor.from_string("FFFFFF")
    bookmark(p, f"part_{number}")

    if subtitle:
        p2 = c.add_paragraph()
        spacing(p2, before=4, after=0)
        r3 = p2.add_run(subtitle)
        r3.font.name = FONT_HEAD
        r3.font.size = Pt(9.5)
        r3.font.italic = True
        r3.font.color.rgb = RGBColor.from_string("D6D6D6")

    tail = doc.add_paragraph()
    spacing(tail, before=0, after=8)
    tail.paragraph_format.line_spacing = 1
    return t


def chapter_head(doc, number, title, aka=None, bm=None):
    """Numbered chapter heading (Heading 2) with a rule beneath."""
    p = doc.add_paragraph(style="Heading 2")
    spacing(p, before=16, after=2)
    keep_together(p)
    r = p.add_run(f"{number}   ")
    r.font.name = FONT_HEAD
    r.font.size = Pt(14.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("666666")
    r2 = p.add_run(title)
    r2.font.name = FONT_HEAD
    r2.font.size = Pt(14.5)
    r2.font.bold = True
    bookmark(p, bm or f"ch_{number}".replace(".", "_"))
    set_borders(p, {"bottom": {"sz": 12, "color": "000000", "space": 3}})

    if aka:
        q = doc.add_paragraph()
        spacing(q, before=3, after=5)
        r3 = q.add_run("Also called:  ")
        r3.font.name = FONT_HEAD
        r3.font.size = Pt(8.6)
        r3.font.bold = True
        r3.font.all_caps = True
        r3.font.color.rgb = RGBColor.from_string("555555")
        r4 = q.add_run(aka)
        r4.font.size = Pt(9.4)
        r4.italic = True
        r4.font.color.rgb = RGBColor.from_string("333333")
    return p


def sub_head(doc, title, bm=None, level=3):
    p = doc.add_paragraph(style=f"Heading {level}")
    keep_together(p)
    r = p.add_run(title)
    r.font.name = FONT_HEAD
    r.bold = True
    if bm:
        bookmark(p, bm)
    return p


def mini_head(doc, title):
    """Small all-caps kicker label above a block."""
    p = doc.add_paragraph(style="TinyLabel")
    p.add_run(title)
    return p


def box(doc, title, lines, fill=G_BOX, accent=True, width=None,
        title_style="BoxHead", body_style="BoxBody", bullets=True,
        marker="▪"):
    """A shaded, ruled callout box (key points / mnemonics / warnings)."""
    width = width or BODY_W
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    c.width = width
    shade(c, fill)
    cell_margins(c, top=110, bottom=110, start=150, end=150)
    edges = {"top": {"sz": 6, "color": "000000"},
             "bottom": {"sz": 6, "color": "000000"},
             "right": {"sz": 6, "color": "000000"},
             "left": {"sz": 24 if accent else 6, "color": "000000"}}
    set_borders(t, edges)

    first = c.paragraphs[0]
    if title:
        first.style = doc.styles[title_style]
        spacing(first, before=0, after=3)
        first.add_run(title)
    else:
        first._p.getparent().remove(first._p)

    for i, line in enumerate(lines):
        p = c.add_paragraph(style=body_style)
        spacing(p, before=1, after=2)
        if bullets:
            indent(p, left=0.36, hanging=0.36)
            r = p.add_run(marker + "  ")
            r.font.name = FONT_HEAD
            r.bold = True
        _rich(p, line)

    tail = doc.add_paragraph()
    spacing(tail, before=0, after=7)
    tail.paragraph_format.line_spacing = 1
    return t


def two_col(doc, left_w=None, right_w=None, gap=Cm(0.35)):
    """
    Layout table: text on the left, figure on the right.
    Returns (left_cell, right_cell). Borderless.
    """
    right_w = right_w or Cm(6.4)
    left_w = left_w or (BODY_W - right_w - gap)
    t = doc.add_table(rows=1, cols=3)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    widths = [left_w, gap, right_w]
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            cell_margins(cell, top=0, bottom=0, start=0, end=0)
    # lock the layout
    tblPr = t._tbl.tblPr
    _sub(tblPr, "w:tblLayout", **{"w:type": "fixed"})
    left, _, right = t.rows[0].cells
    # remove the seed empty paragraphs
    for c in (left, right):
        if c.paragraphs and not c.paragraphs[0].runs:
            c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
    return left, right


def cell_para(cell, text="", style="Normal", before=None, after=None):
    p = cell.add_paragraph(style=style)
    if text:
        _rich(p, text)
    if before is not None or after is not None:
        spacing(p, before=before, after=after)
    return p


def cell_mini_head(cell, title):
    p = cell.add_paragraph(style="TinyLabel")
    p.add_run(title)
    return p


def cell_bullet(cell, text, level=0, marker=None, style=None):
    style = style or ("Bullet" if level == 0 else "SubBullet")
    mk = marker if marker is not None else ("▪" if level == 0 else "–")
    p = cell.add_paragraph(style=style)
    indent(p, left=0.42 + 0.42 * level, hanging=0.42)
    r = p.add_run(mk + "  ")
    r.font.name = FONT_HEAD
    r.bold = level == 0
    _rich(p, text)
    return p


def cell_box(cell, title, lines, fill=G_BOX, marker="▪"):
    """A small shaded callout inside a layout cell (e.g. the side column)."""
    t = cell.add_table(rows=1, cols=1)
    t.autofit = False
    c = t.rows[0].cells[0]
    c.width = cell.width
    shade(c, fill)
    cell_margins(c, top=90, bottom=90, start=120, end=120)
    set_borders(t, {"top": {"sz": 6, "color": "000000"},
                    "bottom": {"sz": 6, "color": "000000"},
                    "right": {"sz": 6, "color": "000000"},
                    "left": {"sz": 20, "color": "000000"}})
    first = c.paragraphs[0]
    if title:
        first.style = "BoxHead"
        spacing(first, before=0, after=3)
        first.add_run(title)
    else:
        first._p.getparent().remove(first._p)
    for line in lines:
        p = c.add_paragraph(style="BoxBody")
        spacing(p, before=1, after=2)
        indent(p, left=0.32, hanging=0.32)
        r = p.add_run(marker + "  ")
        r.font.name = FONT_HEAD
        r.bold = True
        _rich(p, line)
    return t


# kept for backwards compatibility
_cell_para = cell_para


def figure(container, path, width_cm, caption=None, doc=None, border=True,
           align=WD_ALIGN_PARAGRAPH.CENTER, after=6):
    """
    Place an image (optionally inside a table cell) with a thin border and
    an italic caption beneath it.
    """
    adder = container.add_paragraph
    p = adder()
    p.alignment = align
    spacing(p, before=2, after=1)
    p.paragraph_format.line_spacing = 1
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    if border:
        set_borders(p, {"top": {"sz": 4, "color": "999999", "space": 2},
                        "bottom": {"sz": 4, "color": "999999", "space": 2},
                        "left": {"sz": 4, "color": "999999", "space": 3},
                        "right": {"sz": 4, "color": "999999", "space": 3}})
    if caption:
        cp = adder()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.style = _safe_style(container, "FigCaption")
        spacing(cp, before=2, after=after)
        _rich(cp, caption)
    return p


def _safe_style(container, name):
    """Resolve a style object from whatever container we were handed."""
    try:
        part = container.part
        return part.document.styles[name]
    except Exception:
        return name


def instrument_table(doc, groups, width=None, qty_w=Cm(1.5),
                     name_w=None, note_w=None, show_group_rows=True,
                     head=("Qty", "Instrument", "Notes / rationale")):
    """
    The main instrument-listing table.

    groups: list of {"group": str, "items": [(qty, name, note), ...]}
    Renders a header row, then for each group a shaded band row followed by
    its item rows, with zebra banding for legibility in monochrome.
    """
    width = width or BODY_W
    name_w = name_w or Cm(7.2)
    note_w = note_w or (width - qty_w - name_w)

    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    _sub(tblPr, "w:tblLayout", **{"w:type": "fixed"})
    set_borders(t, {
        "top": {"sz": 12, "color": "000000"},
        "bottom": {"sz": 12, "color": "000000"},
        "left": {"sz": 6, "color": "000000"},
        "right": {"sz": 6, "color": "000000"},
        "insideH": {"sz": 4, "color": "BBBBBB"},
        "insideV": {"sz": 4, "color": "BBBBBB"},
    })

    widths = [qty_w, name_w, note_w]

    # ---- header
    hdr = t.rows[0]
    _repeat_header(hdr)
    for cell, txt, w in zip(hdr.cells, head, widths):
        cell.width = w
        shade(cell, G_TABHEAD)
        cell_margins(cell, top=50, bottom=50, start=80, end=80)
        p = cell.paragraphs[0]
        p.style = doc.styles["TabHead"]
        p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if txt == head[0]
                       else WD_ALIGN_PARAGRAPH.LEFT)
        p.add_run(txt)
        set_borders(cell, {"bottom": {"sz": 12, "color": "000000"}})

    band = False
    for g in groups:
        if show_group_rows and g.get("group"):
            row = t.add_row()
            _merge_row(row)
            c = row.cells[0]
            shade(c, G_BOX2)
            cell_margins(c, top=48, bottom=48, start=80, end=80)
            p = c.paragraphs[0]
            p.style = doc.styles["TabHead"]
            keep_together(p)
            r = p.add_run(g["group"].upper())
            r.font.name = FONT_HEAD
            r.font.size = Pt(8.8)
            r.bold = True
            rPr = r._r.get_or_add_rPr()
            rPr.append(_el("w:spacing", **{"w:val": 12}))
            band = False

        for item in g["items"]:
            qty, name, note = (list(item) + ["", "", ""])[:3]
            row = t.add_row()
            if band:
                for c in row.cells:
                    shade(c, G_BAND)
            band = not band
            for cell, txt, w, st, al in zip(
                    row.cells, (qty, name, note), widths,
                    ("TabCell", "TabCell", "TabCellSm"),
                    (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                     WD_ALIGN_PARAGRAPH.LEFT)):
                cell.width = w
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                cell_margins(cell, top=36, bottom=36, start=80, end=80)
                p = cell.paragraphs[0]
                p.style = doc.styles[st]
                p.alignment = al
                if st == "TabCell" and txt is qty:
                    pass
                _rich(p, str(txt))
            # bold the instrument name
            for r in row.cells[1].paragraphs[0].runs:
                r.bold = True

    tail = doc.add_paragraph()
    spacing(tail, before=0, after=8)
    tail.paragraph_format.line_spacing = 1
    return t


def simple_table(doc, head, rows, widths=None, width=None, styles=None,
                 align=None, band=True, first_bold=True, font_size=9.3):
    """A generic comparison / summary table."""
    width = width or BODY_W
    ncol = len(head)
    if widths is None:
        widths = [width / ncol] * ncol
    t = doc.add_table(rows=1, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _sub(t._tbl.tblPr, "w:tblLayout", **{"w:type": "fixed"})
    set_borders(t, {
        "top": {"sz": 12, "color": "000000"},
        "bottom": {"sz": 12, "color": "000000"},
        "left": {"sz": 6, "color": "000000"},
        "right": {"sz": 6, "color": "000000"},
        "insideH": {"sz": 4, "color": "BBBBBB"},
        "insideV": {"sz": 4, "color": "BBBBBB"},
    })
    hdr = t.rows[0]
    _repeat_header(hdr)
    for cell, txt, w in zip(hdr.cells, head, widths):
        cell.width = w
        shade(cell, G_TABHEAD)
        cell_margins(cell, top=50, bottom=50, start=80, end=80)
        p = cell.paragraphs[0]
        p.style = doc.styles["TabHead"]
        p.add_run(txt)
        set_borders(cell, {"bottom": {"sz": 12, "color": "000000"}})

    for i, r in enumerate(rows):
        row = t.add_row()
        if band and i % 2:
            for c in row.cells:
                shade(c, G_BAND)
        for j, (cell, txt, w) in enumerate(zip(row.cells, r, widths)):
            cell.width = w
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            cell_margins(cell, top=36, bottom=36, start=80, end=80)
            p = cell.paragraphs[0]
            p.style = doc.styles["TabCell"]
            for run_ in p.runs:
                run_.font.size = Pt(font_size)
            if align and align[j]:
                p.alignment = align[j]
            _rich(p, str(txt))
            if j == 0 and first_bold:
                for run_ in p.runs:
                    run_.bold = True
    tail = doc.add_paragraph()
    spacing(tail, before=0, after=8)
    tail.paragraph_format.line_spacing = 1
    return t


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    _sub(trPr, "w:tblHeader")
    _sub(trPr, "w:cantSplit")


def _merge_row(row):
    cells = row.cells
    a = cells[0]
    for b in cells[1:]:
        a = a.merge(b)
    return a


# ------------------------------------------------------------------- cover/TOC
def cover_page(doc, title, subtitle, lines, footer_lines=None):
    for _ in range(2):
        p = doc.add_paragraph()
        spacing(p, after=0)

    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c = t.rows[0].cells[0]
    c.width = BODY_W
    cell_margins(c, top=420, bottom=420, start=260, end=260)
    set_borders(t, {"top": {"sz": 36, "color": "000000"},
                    "bottom": {"sz": 36, "color": "000000"},
                    "left": {"sz": 8, "color": "000000"},
                    "right": {"sz": 8, "color": "000000"}})

    p = c.paragraphs[0]
    p.style = doc.styles["CoverSub"]
    spacing(p, after=6)
    r = p.add_run(subtitle.upper())
    r.font.size = Pt(11)
    r.font.bold = True
    rPr = r._r.get_or_add_rPr()
    rPr.append(_el("w:spacing", **{"w:val": 60}))

    p2 = c.add_paragraph(style="CoverTitle")
    spacing(p2, before=2, after=8)
    p2.add_run(title)

    p3 = c.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(p3, before=0, after=10)
    p3.paragraph_format.line_spacing = 1
    set_borders(p3, {"bottom": {"sz": 12, "color": "000000", "space": 1}})

    for ln in lines:
        q = c.add_paragraph(style="CoverSub")
        spacing(q, before=1, after=1)
        rr = q.add_run(ln)
        rr.font.size = Pt(10.5)

    if footer_lines:
        q = c.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(q, before=14, after=2)
        q.paragraph_format.line_spacing = 1
        set_borders(q, {"top": {"sz": 6, "color": "999999", "space": 4}})
        for ln in footer_lines:
            f = c.add_paragraph(style="CoverSub")
            spacing(f, before=1, after=1)
            fr = f.add_run(ln)
            fr.font.size = Pt(9)
            fr.font.italic = True
            fr.font.color.rgb = RGBColor.from_string("555555")
    return t


def toc(doc, title="Table of Contents", levels="1-3",
        note=None):
    p = doc.add_paragraph(style="TOCHead")
    spacing(p, before=6, after=2)
    p.add_run(title)
    bookmark(p, "toc")

    r = doc.add_paragraph()
    r.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(r, before=0, after=4)
    r.paragraph_format.line_spacing = 1
    set_borders(r, {"bottom": {"sz": 18, "color": "000000", "space": 1}})

    if note:
        n = doc.add_paragraph()
        n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(n, before=0, after=10)
        nr = n.add_run(note)
        nr.font.name = FONT_HEAD
        nr.font.size = Pt(8.4)
        nr.font.italic = True
        nr.font.color.rgb = RGBColor.from_string("555555")

    tp = doc.add_paragraph()
    spacing(tp, before=0, after=0)
    add_field(tp,
              f' TOC \\o "{levels}" \\h \\z \\u ',
              placeholder="Right-click here and choose “Update Field”, "
                          "or press Ctrl+A then F9, to build the "
                          "Table of Contents with page numbers.")
    return tp
