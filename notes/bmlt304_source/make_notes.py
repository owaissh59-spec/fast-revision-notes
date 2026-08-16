# -*- coding: utf-8 -*-
"""
make_notes.py
Assembles the complete Medical Ethics (BMLT304) notes into a single A4 DOCX.
"""

import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import builder
from builder import (new_document, build_header_footer, render, page_break,
                     configure_heading_styles, enable_update_fields)
import frontmatter

import ch01_02, ch03, ch04, ch05, ch06, ch07, ch08, ch09, ch10, ch11, appendix

OUT = "/projects/sandbox/Medical_Ethics_BMLT304_Complete_Notes.docx"


def main():
    doc = new_document()
    configure_heading_styles(doc)
    build_header_footer(
        doc,
        "Medical Ethics (BMLT304)  \u2022  B.Sc. MLT 3rd Year  \u2022  Complete Theory Notes",
        "Professionalism and Medical Ethics \u2014 Exam Edition",
    )

    # ---------- front matter ----------
    frontmatter.cover(doc)
    frontmatter.how_to_use(doc)
    frontmatter.contents(doc)
    frontmatter.syllabus_map(doc)

    # ---------- chapters ----------
    modules = [ch01_02, ch03, ch04, ch05, ch06, ch07, ch08, ch09, ch10, ch11, appendix]
    for m in modules:
        render(doc, m.BLOCKS)

    # ---------- closing page ----------
    page_break(doc)
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from builder import shade_cell, cell_borders, cell_margins, HEAD_FONT, BODY_FONT, NAVY, STEEL

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.rows[0].cells[0]
    c.width = Cm(17.2)
    shade_cell(c, NAVY)
    cell_borders(c, color=NAVY, sz=4)
    cell_margins(c, 300, 200, 300, 200)

    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("END OF NOTES")
    r.bold = True; r.font.size = Pt(26); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string("FFFFFF")

    p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Medical Ethics  \u2022  BMLT304  \u2022  All 11 Syllabus Units Covered in Full")
    r.font.size = Pt(11.5); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string("9CC3E5")

    p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("\u201CThe secret of getting ahead is getting started \u2014 and the secret of getting\nfull marks is a definition, a numbered list, a table, and a conclusion.\u201D")
    r.italic = True; r.font.size = Pt(11); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string("F2C14E")

    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(14)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Revise the \u23F1 60-SECOND RECAP boxes and the \u2605 EXAM FOCUS boxes on the morning of the paper.")
    r.bold = True; r.font.size = Pt(10.5); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string(STEEL)

    # ---------- document properties ----------
    cp = doc.core_properties
    cp.title = "Medical Ethics (BMLT304) \u2014 Complete Theory Notes"
    cp.subject = "Professionalism and Medical Ethics \u2014 B.Sc. Medical Laboratory Technology, 3rd Year"
    cp.author = "Exam Edition Study Notes"
    cp.keywords = ("medical ethics; BMLT304; MLT; professionalism; autonomy; beneficence; "
                   "non-maleficence; justice; confidentiality; negligence; consumer protection; "
                   "time management; stress management; leadership; teamwork")
    cp.comments = "Covers all 11 prescribed syllabus topics in full, with definitions, tables, mnemonics and question banks."

    # ---------- make Word refresh the TOC / PAGEREF fields on open ----------
    enable_update_fields(doc)

    doc.save(OUT)
    print("Saved:", OUT)

    # quick stats
    words = 0
    for para in doc.paragraphs:
        words += len(para.text.split())
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    words += len(para.text.split())
    print("Paragraphs:", len(doc.paragraphs))
    print("Tables:", len(doc.tables))
    print("Approx words:", words)


if __name__ == "__main__":
    main()
