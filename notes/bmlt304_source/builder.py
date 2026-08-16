"""
builder.py
Rendering engine for producing attractive, topper-style A4 study notes in DOCX.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ----------------------------------------------------------------------------
# PALETTE
# ----------------------------------------------------------------------------
NAVY = "1F3864"        # deep indigo - chapter bars
BLUE = "1F4E79"        # H2 text
STEEL = "2E74B5"       # rules / accents
ORANGE = "C55A11"      # H3 text
MAROON = "9E2B25"      # warnings
GREEN = "2F6B3A"       # key points
PURPLE = "5B2C87"      # mnemonics
TEAL = "156B6B"        # definitions

FILL_DEF = "FFF6D9"     # definition box
FILL_KEY = "E7F2E4"     # key points box
FILL_TIP = "DEEAF6"     # exam tip box
FILL_MNE = "EDE4F6"     # mnemonic box
FILL_WARN = "FBE2E2"    # warning box
FILL_CASE = "FFF0E1"    # case-law / example box
FILL_TBLHDR = NAVY
FILL_TBLALT = "F3F6FA"
FILL_QBOX = "F0F4F9"

BODY_FONT = "Calibri"
HEAD_FONT = "Cambria"

BODY_SIZE = Pt(10.5)


# ----------------------------------------------------------------------------
# LOW LEVEL XML HELPERS
# ----------------------------------------------------------------------------
def _shade(element, fill):
    """Apply solid background shading to a paragraph or table-cell element."""
    pr = element.get_or_add_pPr() if element.tag.endswith('}p') else element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pr.append(shd)


def shade_paragraph(par, fill):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def paragraph_borders(par, **kwargs):
    """
    kwargs like left=('single', 18, NAVY) -> (style, size(1/8 pt), color)
    sides: top, bottom, left, right
    """
    pPr = par._p.get_or_add_pPr()
    pbdr = pPr.find(qn('w:pBdr'))
    if pbdr is None:
        pbdr = OxmlElement('w:pBdr')
        pPr.append(pbdr)
    for side in ('top', 'left', 'bottom', 'right'):
        if side in kwargs and kwargs[side]:
            style, sz, color = kwargs[side]
            el = OxmlElement('w:%s' % side)
            el.set(qn('w:val'), style)
            el.set(qn('w:sz'), str(sz))
            el.set(qn('w:space'), '4')
            el.set(qn('w:color'), color)
            pbdr.append(el)


def cell_borders(cell, color="BFC9D6", sz=6, sides=('top', 'left', 'bottom', 'right')):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for side in sides:
        el = OxmlElement('w:%s' % side)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)


def cell_margins(cell, top=90, start=140, bottom=90, end=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        el = OxmlElement('w:%s' % tag)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tcPr.append(mar)


def keep_with_next(par):
    pPr = par._p.get_or_add_pPr()
    el = OxmlElement('w:keepNext')
    pPr.append(el)


def no_split(par):
    pPr = par._p.get_or_add_pPr()
    el = OxmlElement('w:keepLines')
    pPr.append(el)


def add_field(par, instr):
    """Insert a Word field code (used for PAGE / NUMPAGES)."""
    run = par.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    return run


# ----------------------------------------------------------------------------
# HEADING STYLES, BOOKMARKS, TOC AND PAGEREF FIELDS
# ----------------------------------------------------------------------------
_BM_ID = [1000]          # bookmark id counter


def configure_heading_styles(doc):
    """Make the built-in Heading 1-4 styles match the visual design.

    Using the *built-in* styles (styleId Heading1..Heading4) is what makes the
    document appear in Word's Navigation Pane, become PDF bookmarks on export,
    and be collected by a TOC field.
    """
    spec = [
        ('Heading 1', 15, NAVY, True, False),
        ('Heading 2', 13, BLUE, True, False),
        ('Heading 3', 11.5, ORANGE, True, False),
        ('Heading 4', 10.5, TEAL, True, True),
    ]
    for name, size, color, bold, italic in spec:
        st = doc.styles[name]
        st.font.name = HEAD_FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = italic
        st.font.color.rgb = RGBColor.from_string(color)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), HEAD_FONT)
        # strip the theme-font references so our font actually wins
        rf = st.element.rPr.rFonts
        for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(attr)) is not None:
                del rf.attrib[qn(attr)]
        rf.set(qn('w:ascii'), HEAD_FONT)
        rf.set(qn('w:hAnsi'), HEAD_FONT)
        pf = st.paragraph_format
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.keep_with_next = True


def add_bookmark(paragraph, name):
    """Wrap a paragraph in a named Word bookmark (usable by PAGEREF / hyperlinks)."""
    _BM_ID[0] += 1
    bid = str(_BM_ID[0])
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), bid)
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)
    return name


def add_pageref(paragraph, bookmark, size=Pt(9.5), color=NAVY, bold=True):
    """Insert a { PAGEREF <bookmark> \\h } field - shows the page number of that bookmark."""
    run = paragraph.add_run()
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), 'begin')
    fc.set(qn('w:dirty'), 'true')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' PAGEREF %s \\h ' % bookmark
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t')
    txt.text = '\u2013'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for el in (fc, it, sep, txt, end):
        run._r.append(el)
    run.font.size = size
    run.font.name = BODY_FONT
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_toc(doc, levels="1-3"):
    """Insert a real { TOC } field which Word fills in with page numbers.

    The w:dirty flag makes Word refresh it as soon as the document is opened,
    and it can be rebuilt any time with right-click > Update Field, or Ctrl+A F9.
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)

    run = p.add_run()
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), 'begin')
    fc.set(qn('w:dirty'), 'true')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' TOC \\o "%s" \\h \\z \\u ' % levels
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    for el in (fc, it, sep):
        run._r.append(el)

    ph = p.add_run("The table of contents will appear here with page numbers. "
                   "If it is still blank, right-click on this line and choose "
                   "\u201CUpdate Field\u201D.")
    ph.italic = True
    ph.font.size = Pt(10)
    ph.font.name = BODY_FONT
    ph.font.color.rgb = RGBColor.from_string("808A99")

    endrun = p.add_run()
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    endrun._r.append(end)
    return p


def enable_update_fields(doc):
    """Ask Word to update all fields (TOC + PAGEREF) when the document opens."""
    settings = doc.settings.element
    for existing in settings.findall(qn('w:updateFields')):
        settings.remove(existing)
    uf = OxmlElement('w:updateFields')
    uf.set(qn('w:val'), 'true')
    # schema order: w:updateFields must precede w:compat / w:rsids
    anchor = settings.find(qn('w:compat'))
    if anchor is None:
        anchor = settings.find(qn('w:rsids'))
    if anchor is not None:
        anchor.addprevious(uf)
    else:
        settings.append(uf)


def horizontal_rule(doc, color=STEEL, sz=12, space_before=4, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    paragraph_borders(p, bottom=('single', sz, color))
    return p


# ----------------------------------------------------------------------------
# RICH TEXT: supports **bold**, *italic*, __underline__, `highlight`
# ----------------------------------------------------------------------------
import re
_TOKEN = re.compile(r"(\*\*.+?\*\*|__.+?__|`.+?`|\*[^*]+?\*)", re.S)


def write_rich(par, text, size=BODY_SIZE, color=None, font=BODY_FONT, bold_color=None):
    """Write text into paragraph honouring lightweight markdown-ish markers."""
    for piece in _TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**') and len(piece) > 4:
            r = par.add_run(piece[2:-2]); r.bold = True
            if bold_color:
                r.font.color.rgb = RGBColor.from_string(bold_color)
            elif color:
                r.font.color.rgb = RGBColor.from_string(color)
        elif piece.startswith('__') and piece.endswith('__') and len(piece) > 4:
            r = par.add_run(piece[2:-2]); r.underline = True
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        elif piece.startswith('`') and piece.endswith('`') and len(piece) > 2:
            r = par.add_run(piece[1:-1]); r.bold = True
            r.font.color.rgb = RGBColor.from_string(MAROON)
        elif piece.startswith('*') and piece.endswith('*') and len(piece) > 2:
            r = par.add_run(piece[1:-1]); r.italic = True
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        else:
            r = par.add_run(piece)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        r.font.size = size
        r.font.name = font
        r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    return par


# ----------------------------------------------------------------------------
# DOCUMENT SET-UP
# ----------------------------------------------------------------------------
def new_document():
    doc = Document()

    # ---- base style
    st = doc.styles['Normal']
    st.font.name = BODY_FONT
    st.font.size = BODY_SIZE
    st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.12

    # ---- A4 page, comfortable margins
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.9)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.8)
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(1.0)

    # ---- list styles must exist for numbering to look right
    return doc


def build_header_footer(doc, header_text, footer_left):
    sec = doc.sections[0]

    hdr = sec.header
    hp = hdr.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(header_text)
    r.font.size = Pt(8)
    r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string("7F8CA0")
    r.italic = True
    paragraph_borders(hp, bottom=('single', 6, "C6D0DE"))

    ftr = sec.footer
    fp = ftr.paragraphs[0]
    fp.text = ""
    paragraph_borders(fp, top=('single', 6, "C6D0DE"))
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.2), WD_ALIGN_PARAGRAPH.RIGHT)
    r = fp.add_run(footer_left + "\t")
    r.font.size = Pt(8); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string("7F8CA0")
    r2 = fp.add_run("Page ")
    r2.font.size = Pt(8); r2.font.name = BODY_FONT
    r2.font.color.rgb = RGBColor.from_string(NAVY); r2.bold = True
    fr = add_field(fp, "PAGE")
    fr.font.size = Pt(8); fr.font.name = BODY_FONT
    fr.font.color.rgb = RGBColor.from_string(NAVY); fr.bold = True
    r3 = fp.add_run(" of ")
    r3.font.size = Pt(8); r3.font.name = BODY_FONT
    r3.font.color.rgb = RGBColor.from_string("7F8CA0")
    fr2 = add_field(fp, "NUMPAGES")
    fr2.font.size = Pt(8); fr2.font.name = BODY_FONT
    fr2.font.color.rgb = RGBColor.from_string("7F8CA0")


# ----------------------------------------------------------------------------
# BLOCK RENDERERS
# ----------------------------------------------------------------------------
def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def chapter_bar(doc, number, title, subtitle=None, bookmark=None):
    """Full-width navy chapter banner.

    The main line is a genuine **Heading 1** paragraph, so the chapter shows up
    in Word's Navigation Pane, becomes a PDF bookmark on export, and is picked
    up by the TOC field. It is also wrapped in a named bookmark so its page
    number can be referenced with PAGEREF.
    """
    label = "Appendix" if str(number).upper() == "A" else "Chapter %s" % number

    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(0.55), Cm(16.65)]
    c0, c1 = tbl.rows[0].cells
    for c, w in zip((c0, c1), widths):
        c.width = w
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)

    shade_cell(c0, ORANGE)
    shade_cell(c1, NAVY)
    cell_borders(c0, color=ORANGE, sz=4)
    cell_borders(c1, color=NAVY, sz=4)
    cell_margins(c0, 0, 0, 0, 0)
    cell_margins(c1, 160, 200, 160, 140)

    # ---- main heading line (Heading 1) ----
    p1 = c1.paragraphs[0]
    p1.style = doc.styles['Heading 1']
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    p1.paragraph_format.line_spacing = 1.0

    r = p1.add_run(label)
    r.bold = True; r.font.size = Pt(11); r.font.name = HEAD_FONT
    r.font.all_caps = True
    r.font.color.rgb = RGBColor.from_string("F2C14E")

    r = p1.add_run("   \u2022   ")
    r.bold = True; r.font.size = Pt(11); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string("7FA6D0")

    r = p1.add_run(title)
    r.bold = True; r.font.size = Pt(15); r.font.name = HEAD_FONT
    r.font.all_caps = True
    r.font.color.rgb = RGBColor.from_string("FFFFFF")

    if bookmark:
        add_bookmark(p1, bookmark)

    if subtitle:
        p2 = c1.add_paragraph()
        p2.paragraph_format.space_before = Pt(3)
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(subtitle)
        r.italic = True; r.font.size = Pt(9.5); r.font.name = BODY_FONT
        r.font.color.rgb = RGBColor.from_string("CBD8EA")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    spacer.paragraph_format.space_before = Pt(0)
    return tbl


def h2(doc, text):
    """Section heading - a genuine Heading 2 (navigable, TOC level 2)."""
    p = doc.add_paragraph(style=doc.styles['Heading 2'])
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Cm(0.0)
    keep_with_next(p)
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(13); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph_borders(p, bottom=('single', 10, "9DB7D4"),
                      left=('single', 24, ORANGE))
    p.paragraph_format.left_indent = Cm(0.22)
    return p


def h3(doc, text):
    """Sub-section heading - a genuine Heading 3 (navigable, TOC level 3)."""
    p = doc.add_paragraph(style=doc.styles['Heading 3'])
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.32)
    keep_with_next(p)
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(11.5); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string(ORANGE)
    paragraph_borders(p, left=('single', 12, "E8B27A"))
    return p


def h4(doc, text):
    """Minor heading - a genuine Heading 4 (navigable, excluded from the TOC)."""
    p = doc.add_paragraph(style=doc.styles['Heading 4'])
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.4)
    keep_with_next(p)
    r = p.add_run(text)
    r.bold = True; r.italic = True; r.font.size = Pt(10.5); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string(TEAL)
    return p


def para(doc, text, indent=0.2, justify=True):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(5)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    write_rich(p, text, bold_color=NAVY)
    return p


def bullets(doc, items, indent=0.55, marker="\u25CF", marker_color=STEEL, size=BODY_SIZE):
    out = []
    for it in items:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(indent + 0.45)
        pf.first_line_indent = Cm(-0.45)
        pf.space_after = Pt(2.5)
        pf.line_spacing = 1.10
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(marker + "  ")
        r.font.size = Pt(7.5 if marker == "\u25CF" else 9)
        r.font.color.rgb = RGBColor.from_string(marker_color)
        r.bold = True
        write_rich(p, it, size=size, bold_color=NAVY)
        out.append(p)
    return out


def subbullets(doc, items):
    return bullets(doc, items, indent=1.15, marker="\u2013", marker_color=ORANGE,
                   size=Pt(10))


def numbers(doc, items, indent=0.55, start=1):
    out = []
    for i, it in enumerate(items, start):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(indent + 0.62)
        pf.first_line_indent = Cm(-0.62)
        pf.space_after = Pt(2.5)
        pf.line_spacing = 1.10
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run("%d." % i)
        r.bold = True; r.font.size = BODY_SIZE
        r.font.color.rgb = RGBColor.from_string(NAVY)
        r2 = p.add_run("  ")
        r2.font.size = BODY_SIZE
        write_rich(p, it, bold_color=NAVY)
        out.append(p)
    return out


def _box(doc, label, body_lines, fill, accent, label_color=None, icon=""):
    """Generic single-cell coloured callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Cm(17.2)
    shade_cell(cell, fill)
    cell_borders(cell, color=accent, sz=4)
    # thick left accent bar
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    left = borders.find(qn('w:left'))
    left.set(qn('w:sz'), '24')
    cell_margins(cell, 110, 170, 110, 150)

    first = cell.paragraphs[0]
    first.paragraph_format.space_after = Pt(3)
    first.paragraph_format.space_before = Pt(0)
    if label:
        r = first.add_run(("%s %s" % (icon, label)).strip())
        r.bold = True
        r.font.size = Pt(10)
        r.font.name = BODY_FONT
        r.font.color.rgb = RGBColor.from_string(label_color or accent)
        target = None
    else:
        target = first

    for i, line in enumerate(body_lines):
        if target is not None and i == 0:
            p = target
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.10
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if line.startswith("- "):
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            rr = p.add_run("\u25AA  ")
            rr.font.size = Pt(8); rr.bold = True
            rr.font.color.rgb = RGBColor.from_string(accent)
            write_rich(p, line[2:], size=Pt(10), bold_color=accent)
        else:
            write_rich(p, line, size=Pt(10), bold_color=accent)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.space_before = Pt(0)
    return tbl


def definition(doc, term, text, extra=None):
    lines = ["**%s:** %s" % (term, text)]
    if extra:
        lines += extra
    return _box(doc, "DEFINITION", lines, FILL_DEF, "B8860B", label_color="8A6100",
                icon="\u270D")


def keypoints(doc, items, label="KEY POINTS TO REMEMBER"):
    return _box(doc, label, ["- " + i for i in items], FILL_KEY, GREEN, icon="\u2714")


def examtip(doc, lines, label="EXAM FOCUS"):
    if isinstance(lines, str):
        lines = [lines]
    return _box(doc, label, lines, FILL_TIP, BLUE, icon="\u2605")


def mnemonic(doc, name, lines, label=None):
    body = ["**%s**" % name] + lines
    return _box(doc, label or "MEMORY TRICK / MNEMONIC", body, FILL_MNE, PURPLE,
                icon="\U0001F9E0")


def warning(doc, lines, label="IMPORTANT \u2013 DO NOT CONFUSE"):
    if isinstance(lines, str):
        lines = [lines]
    return _box(doc, label, lines, FILL_WARN, MAROON, icon="\u26A0")


def example(doc, lines, label="EXAMPLE / CASE"):
    if isinstance(lines, str):
        lines = [lines]
    return _box(doc, label, lines, FILL_CASE, "A0522D", icon="\U0001F4CE")


def quote(doc, text, source=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.right_indent = Cm(1.0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_borders(p, left=('single', 18, ORANGE), right=('single', 18, ORANGE))
    r = p.add_run('\u201C%s\u201D' % text)
    r.italic = True; r.font.size = Pt(10.5); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string(NAVY)
    if source:
        r2 = p.add_run("  \u2014 %s" % source)
        r2.font.size = Pt(9); r2.bold = True; r2.font.name = BODY_FONT
        r2.font.color.rgb = RGBColor.from_string(ORANGE)
    return p


def table(doc, headers, rows, widths=None, caption=None, header_fill=FILL_TBLHDR,
          font_size=9.5):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(8)
        cp.paragraph_format.space_after = Pt(3)
        cp.paragraph_format.left_indent = Cm(0.2)
        keep_with_next(cp)
        r = cp.add_run("\u25A6  " + caption)
        r.bold = True; r.font.size = Pt(10); r.font.name = BODY_FONT
        r.font.color.rgb = RGBColor.from_string(NAVY)

    ncols = len(headers)
    tbl = doc.add_table(rows=1, cols=ncols)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    total = 17.2
    if widths is None:
        widths = [total / ncols] * ncols
    else:
        s = float(sum(widths))
        widths = [w / s * total for w in widths]

    # header
    hdr = tbl.rows[0]
    _repeat_header(hdr)
    for i, htext in enumerate(headers):
        c = hdr.cells[i]
        c.width = Cm(widths[i])
        shade_cell(c, header_fill)
        cell_borders(c, color=header_fill, sz=6)
        cell_margins(c, 80, 110, 80, 110)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        clean = htext.replace('**', '').replace('__', '').replace('`', '')
        r = p.add_run(clean)
        r.bold = True; r.font.size = Pt(9.5); r.font.name = BODY_FONT
        r.font.color.rgb = RGBColor.from_string("FFFFFF")

    # body
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci, val in enumerate(row):
            c = cells[ci]
            c.width = Cm(widths[ci])
            cell_borders(c, color="C3CEDD", sz=4)
            cell_margins(c, 70, 110, 70, 110)
            if ri % 2 == 1:
                shade_cell(c, FILL_TBLALT)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.06
            segs = str(val).split("\n")
            for si, seg in enumerate(segs):
                pp = p if si == 0 else c.add_paragraph()
                pp.paragraph_format.space_after = Pt(0)
                pp.paragraph_format.space_before = Pt(0)
                pp.paragraph_format.line_spacing = 1.06
                if seg.startswith("- "):
                    pp.paragraph_format.left_indent = Cm(0.32)
                    pp.paragraph_format.first_line_indent = Cm(-0.32)
                    rr = pp.add_run("\u2022 ")
                    rr.font.size = Pt(font_size)
                    rr.font.color.rgb = RGBColor.from_string(STEEL)
                    seg = seg[2:]
                write_rich(pp, seg, size=Pt(font_size), bold_color=NAVY)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(5)
    sp.paragraph_format.space_before = Pt(0)
    return tbl


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), "true")
    trPr.append(el)


def flow(doc, steps, color=STEEL):
    """Vertical arrow flow-chart rendered as centred shaded boxes."""
    for i, s in enumerate(steps):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        c = tbl.rows[0].cells[0]
        c.width = Cm(11.5)
        shade_cell(c, "EAF1F8")
        cell_borders(c, color=color, sz=8)
        cell_margins(c, 70, 130, 70, 130)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        write_rich(p, s, size=Pt(10), bold_color=NAVY)
        if i != len(steps) - 1:
            ap = doc.add_paragraph()
            ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ap.paragraph_format.space_after = Pt(0)
            ap.paragraph_format.space_before = Pt(0)
            r = ap.add_run("\u2193")
            r.bold = True; r.font.size = Pt(12)
            r.font.color.rgb = RGBColor.from_string(color)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(5)


def qa_section(doc, title, items, marks=None):
    """Exam question bank block."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    shade_paragraph(p, NAVY)
    paragraph_borders(p, left=('single', 20, ORANGE))
    p.paragraph_format.left_indent = Cm(0.15)
    r = p.add_run("  \u2753  " + title)
    r.bold = True; r.font.size = Pt(11); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string("FFFFFF")

    for i, (q, a) in enumerate(items, 1):
        qp = doc.add_paragraph()
        qp.paragraph_format.left_indent = Cm(0.75)
        qp.paragraph_format.first_line_indent = Cm(-0.6)
        qp.paragraph_format.space_before = Pt(5)
        qp.paragraph_format.space_after = Pt(1)
        qp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        keep_with_next(qp)
        r = qp.add_run("Q%d. " % i)
        r.bold = True; r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(MAROON)
        write_rich(qp, q, size=Pt(10), bold_color=NAVY)

        alines = a if isinstance(a, list) else [a]
        for j, al in enumerate(alines):
            ap = doc.add_paragraph()
            ap.paragraph_format.left_indent = Cm(1.35)
            ap.paragraph_format.first_line_indent = Cm(-0.55)
            ap.paragraph_format.space_after = Pt(1)
            ap.paragraph_format.line_spacing = 1.08
            ap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            lab = ap.add_run("Ans. " if j == 0 else "\u25AB  ")
            lab.bold = True; lab.font.size = Pt(9.5)
            lab.font.color.rgb = RGBColor.from_string(GREEN)
            write_rich(ap, al, size=Pt(9.5), bold_color=NAVY)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)


def recap(doc, lines):
    """End-of-chapter revision strip."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.rows[0].cells[0]
    c.width = Cm(17.2)
    shade_cell(c, "FFFBEA")
    cell_borders(c, color=ORANGE, sz=12)
    cell_margins(c, 120, 170, 120, 150)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("\u23F1  LAST-MINUTE REVISION \u2014 60 SECOND RECAP")
    r.bold = True; r.font.size = Pt(10.5); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string("A0522D")
    for line in lines:
        pp = c.add_paragraph()
        pp.paragraph_format.space_after = Pt(2)
        pp.paragraph_format.left_indent = Cm(0.45)
        pp.paragraph_format.first_line_indent = Cm(-0.45)
        pp.paragraph_format.line_spacing = 1.10
        pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rr = pp.add_run("\u27A4  ")
        rr.font.size = Pt(8.5); rr.bold = True
        rr.font.color.rgb = RGBColor.from_string(ORANGE)
        write_rich(pp, line, size=Pt(10), bold_color=MAROON)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    return tbl


# ----------------------------------------------------------------------------
# DISPATCHER
# ----------------------------------------------------------------------------
def render(doc, blocks):
    for block in blocks:
        kind = block[0]
        arg = block[1] if len(block) > 1 else None
        if kind == 'ch':
            page_break(doc)
            num = arg[0]
            try:
                bm = "Ch%02d" % int(num)
            except (TypeError, ValueError):
                bm = "Ch%s" % str(num).upper()
            chapter_bar(doc, num, arg[1],
                        arg[2] if len(arg) > 2 else None, bookmark=bm)
        elif kind == 'h2':
            h2(doc, arg)
        elif kind == 'h3':
            h3(doc, arg)
        elif kind == 'h4':
            h4(doc, arg)
        elif kind == 'p':
            para(doc, arg)
        elif kind == 'b':
            bullets(doc, arg)
        elif kind == 'sb':
            subbullets(doc, arg)
        elif kind == 'n':
            numbers(doc, arg)
        elif kind == 'def':
            definition(doc, arg[0], arg[1], arg[2] if len(arg) > 2 else None)
        elif kind == 'key':
            keypoints(doc, arg)
        elif kind == 'tip':
            examtip(doc, arg)
        elif kind == 'mne':
            mnemonic(doc, arg[0], arg[1])
        elif kind == 'warn':
            warning(doc, arg)
        elif kind == 'eg':
            example(doc, arg)
        elif kind == 'quote':
            quote(doc, arg[0], arg[1] if len(arg) > 1 else None)
        elif kind == 'tbl':
            table(doc, arg[0], arg[1], arg[2] if len(arg) > 2 else None,
                  arg[3] if len(arg) > 3 else None)
        elif kind == 'flow':
            flow(doc, arg)
        elif kind == 'qa':
            qa_section(doc, arg[0], arg[1])
        elif kind == 'recap':
            recap(doc, arg)
        elif kind == 'hr':
            horizontal_rule(doc)
        elif kind == 'pb':
            page_break(doc)
        else:
            raise ValueError("Unknown block kind: %s" % kind)
