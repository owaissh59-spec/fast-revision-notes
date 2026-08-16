"""Cover page, syllabus map and table of contents."""

from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from builder import (NAVY, BLUE, STEEL, ORANGE, MAROON, GREEN, PURPLE, TEAL,
                     HEAD_FONT, BODY_FONT, shade_cell, cell_borders, cell_margins,
                     shade_paragraph, paragraph_borders, page_break, table,
                     write_rich, keypoints, examtip)

CHAPTERS = [
    (1, "Introduction to Professionalism and Ethics",
        "Meaning, characteristics, attributes, code of conduct, ethics vs law vs etiquette"),
    (2, "Value and Dignity of Human Life",
        "Sanctity of life, human dignity, autonomy, respect for persons, end-of-life issues"),
    (3, "Principles of Medical Ethics",
        "Autonomy, Beneficence, Non-maleficence, Justice + supporting principles"),
    (4, "Communication Skills for Interacting with Colleagues, Clinicians, Patients and Attendants",
        "Process, types, barriers, verbal / non-verbal, listening, SBAR, empathy"),
    (5, "Relationship of Paramedics and Patient",
        "Models, rights and duties, informed consent, trust, difficult patients"),
    (6, "Communicating Diagnostic Results and Confidentiality",
        "Critical values, panic value protocol, breaking bad news, HIPAA, exceptions"),
    (7, "Negligence, Malpractice, Legal Implications and Law Suits in Medical Practice",
        "4 D's, res ipsa loquitur, vicarious liability, defences, documentation"),
    (8, "Consumer Protection and Insurance for Professional Health Hazard",
        "CPA 1986 / 2019, redressal machinery, indemnity insurance, biohazards"),
    (9, "Time Management",
        "Principles, Eisenhower matrix, Pareto, time wasters, techniques in the lab"),
    (10, "Stress Management",
        "GAS, eustress vs distress, burnout, coping strategies, relaxation"),
    (11, "Leadership Qualities and Team Work in Health Care Professionals",
        "Styles, theories, qualities, team building, Tuckman, conflict resolution"),
]


def cover(doc):
    # ---------- top colour band ----------
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.rows[0].cells[0]
    c.width = Cm(17.2)
    shade_cell(c, NAVY)
    cell_borders(c, color=NAVY, sz=4)
    cell_margins(c, 260, 200, 260, 200)

    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("B.Sc. MEDICAL LABORATORY TECHNOLOGY  \u2022  3rd YEAR")
    r.bold = True; r.font.size = Pt(11); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string("9CC3E5")

    p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("MEDICAL ETHICS")
    r.bold = True; r.font.size = Pt(40); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string("FFFFFF")

    p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("PROFESSIONALISM  AND  MEDICAL  ETHICS")
    r.bold = True; r.font.size = Pt(13); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string("F2C14E")

    p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Course Code : BMLT304   |   Subsidiary Subject")
    r.bold = True; r.font.size = Pt(12); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string("FFFFFF")

    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(10)

    # ---------- title strip ----------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("COMPLETE THEORY NOTES  \u2014  EXAM EDITION")
    r.bold = True; r.font.size = Pt(16); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string(MAROON)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("All 11 Syllabus Units  \u2022  Definitions  \u2022  Tables  \u2022  Diagrams  \u2022  Mnemonics  \u2022  Expected Questions")
    r.italic = True; r.font.size = Pt(10.5); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string(BLUE)

    # ---------- feature grid ----------
    feats = [
        ("\u2714", "100% Syllabus Coverage",
         "Every one of the 11 prescribed topics is written out in full \u2014 nothing summarised, nothing skipped."),
        ("\u2605", "Exam-Ready Format",
         "Definitions, headed points, comparison tables and numbered lists \u2014 exactly how examiners like answers."),
        ("\U0001F9E0", "Mnemonics & Memory Aids",
         "Built-in memory tricks so long lists can be reproduced under exam pressure."),
        ("\u2753", "Question Bank in Every Chapter",
         "Very-short, short and long answer questions with model answer skeletons."),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i in range(0, len(feats), 2):
        cells = t.add_row().cells
        for j in range(2):
            if i + j >= len(feats):
                continue
            icon, head, body = feats[i + j]
            cc = cells[j]
            cc.width = Cm(8.6)
            shade_cell(cc, "F2F6FB")
            cell_borders(cc, color="AFC6E0", sz=4)
            cell_margins(cc, 130, 150, 130, 150)
            pp = cc.paragraphs[0]
            pp.paragraph_format.space_after = Pt(2)
            rr = pp.add_run("%s  %s" % (icon, head))
            rr.bold = True; rr.font.size = Pt(10.5); rr.font.name = BODY_FONT
            rr.font.color.rgb = RGBColor.from_string(NAVY)
            pp2 = cc.add_paragraph()
            pp2.paragraph_format.space_after = Pt(0)
            pp2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            rr2 = pp2.add_run(body)
            rr2.font.size = Pt(9); rr2.font.name = BODY_FONT
            rr2.font.color.rgb = RGBColor.from_string("404A57")

    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(14)

    # ---------- quote ----------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    paragraph_borders(p, top=('single', 10, ORANGE), bottom=('single', 10, ORANGE))
    r = p.add_run("\u201CThe good physician treats the disease; the great physician treats the patient who has the disease.\u201D")
    r.italic = True; r.font.size = Pt(11); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("\u2014 Sir William Osler")
    r.bold = True; r.font.size = Pt(9.5); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string(ORANGE)

    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(16)

    # ---------- footer band ----------
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.rows[0].cells[0]
    c.width = Cm(17.2)
    shade_cell(c, STEEL)
    cell_borders(c, color=STEEL, sz=4)
    cell_margins(c, 120, 150, 120, 150)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("As per Annexure to Notification No. F (Prescription\u2013Syllabus/Paramedical Courses/Acad/KU/21) dated 23\u201302\u20132021")
    r.bold = True; r.font.size = Pt(9); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string("FFFFFF")


def how_to_use(doc):
    page_break(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    shade_paragraph(p, NAVY)
    p.paragraph_format.left_indent = Cm(0.15)
    r = p.add_run("  HOW TO USE THESE NOTES  \u2014  READ THIS FIRST")
    r.bold = True; r.font.size = Pt(14); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string("FFFFFF")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    write_rich(p, "These notes are written for a **purely theoretical paper**. Medical Ethics is a *scoring* subject because answers are descriptive and the examiner mainly checks whether you have (a) given a **correct definition**, (b) produced **enough headed points**, and (c) added a **relevant example**. The notes below are therefore deliberately written in *point-and-heading* form rather than long paragraphs, so that whatever you read is directly reproducible in the answer book.", bold_color=NAVY)

    from builder import table as _t
    _t(doc,
       ["Symbol / Box", "What it means", "How to use it in the exam"],
       [["\u270D  DEFINITION (yellow box)",
         "The formal, examiner-approved definition of a term.",
         "**Memorise word-for-word.** Always begin a long answer with this. Worth 1\u20132 marks on its own."],
        ["\u2714  KEY POINTS (green box)",
         "The compulsory points of that topic.",
         "These are the **minimum** points needed for full marks. Reproduce all of them."],
        ["\u2605  EXAM FOCUS (blue box)",
         "How the topic is actually asked in the paper.",
         "Tells you whether to expect a 2-mark, 5-mark or 10-mark question."],
        ["\U0001F9E0  MNEMONIC (purple box)",
         "A memory trick for long lists.",
         "Write the mnemonic in the margin of your rough sheet the moment the exam starts."],
        ["\u26A0  IMPORTANT (red box)",
         "Commonly confused pairs / examiner traps.",
         "Most students lose marks here. Read these twice."],
        ["\U0001F4CE  EXAMPLE / CASE (peach box)",
         "A real or illustrative example, or a legal case.",
         "Add one example to every long answer \u2014 it visibly separates a topper's answer."],
        ["\u25A6  TABLE",
         "Comparisons and classifications.",
         "**Draw the table** in the answer sheet. Tables earn marks fast and look organised."],
        ["\u2753  QUESTION BANK",
         "Expected questions with answer skeletons.",
         "Use for final revision \u2014 attempt from memory, then check."],
        ["\u23F1  60-SECOND RECAP",
         "Chapter compressed into a few lines.",
         "Read only these boxes on the morning of the exam."]],
       widths=[3.4, 5.0, 6.4],
       caption="Legend \u2014 the colour-coded boxes used throughout these notes")

    examtip(doc, [
        "**How to structure ANY 10-mark answer in this subject (universal template):**",
        "- **1. Introduction / Definition** \u2014 2\u20133 lines + formal definition.",
        "- **2. Classification / Types / Components** \u2014 in headed points, one line each.",
        "- **3. Detailed explanation** \u2014 expand each point in 2\u20133 lines.",
        "- **4. A table or a simple flow diagram** \u2014 always include at least one.",
        "- **5. Example / case / clinical-laboratory application** \u2014 links theory to MLT practice.",
        "- **6. Conclusion** \u2014 2 lines stating the professional importance.",
        "Following this template alone typically converts a 6/10 answer into a 9\u201310/10 answer."
    ], label="THE UNIVERSAL ANSWER TEMPLATE")

    keypoints(doc, [
        "**Never leave a question blank.** In ethics papers even general professional common-sense written in points earns partial marks.",
        "**Underline every technical term** (autonomy, beneficence, res ipsa loquitur, vicarious liability) with a pen \u2014 examiners scan for keywords.",
        "**Number your points.** Ten numbered points score better than the same content in one paragraph.",
        "**Quote the law correctly** where relevant: Consumer Protection Act **2019** (which replaced the 1986 Act), Indian Medical Council (Professional Conduct, Etiquette and Ethics) Regulations **2002**, Clinical Establishments Act **2010**.",
        "**Draw a box or diagram** for at least two answers in the paper \u2014 visual answers stand out in a bundle of 200 scripts.",
        "**Manage time**: allot roughly 1.5 minutes per mark, and leave 10 minutes at the end for revision \u2014 ironically, this is Chapter 9 of your own syllabus."
    ], label="SIX RULES OF A TOPPER'S ANSWER SHEET")


def _banner(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    shade_paragraph(p, NAVY)
    p.paragraph_format.left_indent = Cm(0.15)
    r = p.add_run("  " + text)
    r.bold = True; r.font.size = Pt(14); r.font.name = HEAD_FONT
    r.font.color.rgb = RGBColor.from_string("FFFFFF")
    return p


def contents(doc):
    """Auto-generating table of contents (real TOC field, with page numbers)."""
    from builder import add_toc

    page_break(doc)
    _banner(doc, "TABLE OF CONTENTS")

    keypoints(doc, [
        "**The page numbers below are live fields.** To refresh them after any edit or reprint: "
        "press **Ctrl + A** (select all) and then **F9** \u2014 that updates this contents list and "
        "the chapter page numbers on the next page in one go. If Word asks, choose "
        "**\u201CUpdate entire table\u201D**.",
        "You can also right-click anywhere on the contents list and choose **\u201CUpdate Field\u201D**.",
        "**To jump to any chapter,** hold **Ctrl** and click its line below \u2014 or open the "
        "**Navigation Pane** (*View \u2192 Navigation Pane*, or **Ctrl + F**) which lists every chapter "
        "and every numbered section as a clickable bookmark.",
        "When you save as PDF, tick **\u201CCreate bookmarks using: Headings\u201D** so the PDF gets a "
        "clickable chapter sidebar too."
    ], label="HOW TO UPDATE THE PAGE NUMBERS \u2014 ONE CLICK")

    add_toc(doc, levels="1-3")


def syllabus_map(doc):
    from builder import add_pageref, cell_borders as _cb, cell_margins as _cm

    page_break(doc)
    _banner(doc, "SYLLABUS MAP  \u2014  CHAPTER FINDER")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    write_rich(p, "**Course Title:** Medical Ethics (Subsidiary Subject)  |  **Course Code:** BMLT304  |  **Class:** B.Sc. 3rd Year Medical Laboratory Technology  |  **Unit heading in syllabus:** *Professionalism and Medical Ethics \u2014 3rd Year*", bold_color=NAVY)

    # ---- manual table so that a live PAGEREF field can go in the last column ----
    headers = ["Ch.", "Topic as printed in the syllabus  \u2014  with the sub-topics covered in these notes", "Page"]
    widths = [1.0, 14.6, 1.6]

    cp = doc.add_paragraph()
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(3)
    cp.paragraph_format.left_indent = Cm(0.2)
    r = cp.add_run("\u25A6  All eleven prescribed topics, with live page numbers")
    r.bold = True; r.font.size = Pt(10); r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor.from_string(NAVY)

    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    hdr = tbl.rows[0]
    for i, htext in enumerate(headers):
        c = hdr.cells[i]
        c.width = Cm(widths[i])
        shade_cell(c, NAVY)
        _cb(c, color=NAVY, sz=6)
        _cm(c, 80, 110, 80, 110)
        pp = c.paragraphs[0]
        pp.paragraph_format.space_after = Pt(0)
        pp.paragraph_format.space_before = Pt(0)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(htext)
        rr.bold = True; rr.font.size = Pt(9.5); rr.font.name = BODY_FONT
        rr.font.color.rgb = RGBColor.from_string("FFFFFF")

    for idx, (num, title, sub) in enumerate(CHAPTERS):
        cells = tbl.add_row().cells
        for ci in range(3):
            cells[ci].width = Cm(widths[ci])
            _cb(cells[ci], color="C3CEDD", sz=4)
            _cm(cells[ci], 70, 110, 70, 110)
            if idx % 2 == 1:
                shade_cell(cells[ci], "F3F6FA")

        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p0.add_run(str(num))
        rr.bold = True; rr.font.size = Pt(10); rr.font.name = BODY_FONT
        rr.font.color.rgb = RGBColor.from_string(NAVY)

        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.06
        write_rich(p1, "**%s**" % title, size=Pt(9.5), bold_color=NAVY)
        p1b = cells[1].add_paragraph()
        p1b.paragraph_format.space_after = Pt(0)
        p1b.paragraph_format.space_before = Pt(0)
        p1b.paragraph_format.line_spacing = 1.06
        write_rich(p1b, "*%s*" % sub, size=Pt(9), bold_color=NAVY)

        p2 = cells[2].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_pageref(p2, "Ch%02d" % num, size=Pt(10))

    # appendix row
    cells = tbl.add_row().cells
    for ci in range(3):
        cells[ci].width = Cm(widths[ci])
        _cb(cells[ci], color="C3CEDD", sz=4)
        _cm(cells[ci], 70, 110, 70, 110)
        shade_cell(cells[ci], "FFF6D9")
    p0 = cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(0)
    rr = p0.add_run("A")
    rr.bold = True; rr.font.size = Pt(10); rr.font.name = BODY_FONT
    rr.font.color.rgb = RGBColor.from_string(NAVY)
    p1 = cells[1].paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    write_rich(p1, "**Appendix \u2014 Rapid Revision Bank**", size=Pt(9.5), bold_color=NAVY)
    p1b = cells[1].add_paragraph()
    p1b.paragraph_format.space_after = Pt(0)
    p1b.paragraph_format.space_before = Pt(0)
    write_rich(p1b, "*Mnemonic index \u2022 legal timeline \u2022 glossary \u2022 answer skeletons \u2022 2-mark definitions*",
               size=Pt(9), bold_color=NAVY)
    p2 = cells[2].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    add_pageref(p2, "ChA", size=Pt(10))

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(5)

    examtip(doc, [
        "**Weightage strategy (based on the way this paper is normally set):**",
        "- **Very high yield (almost certain long question):** Ch. 3 Principles of Medical Ethics, Ch. 7 Negligence & Malpractice, Ch. 4 Communication Skills, Ch. 6 Confidentiality.",
        "- **High yield (short notes / 5-mark):** Ch. 1 Professionalism, Ch. 5 Paramedic\u2013Patient Relationship, Ch. 8 Consumer Protection, Ch. 11 Leadership & Teamwork.",
        "- **Sure-shot short notes:** Ch. 2 Value & Dignity of Human Life, Ch. 9 Time Management, Ch. 10 Stress Management.",
        "Prepare **all eleven** \u2014 the syllabus is small and the paper often picks the 'easy-looking' topics such as Time Management for a full 10-mark question."
    ], label="WHICH CHAPTER CARRIES HOW MUCH WEIGHT")
